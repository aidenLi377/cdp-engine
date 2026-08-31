"""Safe parsing and validation for dimension-table Excel imports."""

from __future__ import annotations

import math
import zipfile
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .constants import REQUIRED_DIMENSION_COLUMNS

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 20_000
MAX_CELL_CHARACTERS = 2_000
MAX_ARCHIVE_FILES = 1_000
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
IDENTIFIER_COLUMNS = {"ID", "Value", "parentId", "BizID", "cateId"}


class DimensionImportFileError(ValueError):
    pass


def _issue(row: int, column: str, code: str, message: str) -> dict:
    return {
        "row": row,
        "column": column,
        "code": code,
        "message": message,
    }


def _validate_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files) > MAX_ARCHIVE_FILES:
                raise DimensionImportFileError("Excel 文件结构过于复杂，请精简后重试")
            if sum(item.file_size for item in files) > MAX_UNCOMPRESSED_BYTES:
                raise DimensionImportFileError("Excel 解压后体积过大，请拆分后重试")
            if "[Content_Types].xml" not in archive.namelist():
                raise DimensionImportFileError("文件不是有效的 Excel 工作簿")
    except zipfile.BadZipFile as exc:
        raise DimensionImportFileError("文件不是有效的 .xlsx 工作簿") from exc


def _cell_text(cell, row_number: int, column: str, issues: list[dict]) -> str:
    if cell.data_type == "f":
        issues.append(
            _issue(row_number, column, "FORMULA_NOT_ALLOWED", "不支持公式，请先粘贴为值")
        )
        return ""
    if cell.data_type == "e":
        issues.append(
            _issue(row_number, column, "CELL_ERROR", "单元格包含 Excel 错误值")
        )
        return ""

    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        issues.append(
            _issue(row_number, column, "BOOLEAN_NOT_ALLOWED", "不支持布尔值，请改为文本")
        )
        return ""
    if isinstance(value, (datetime, date, time)):
        issues.append(
            _issue(row_number, column, "DATE_NOT_ALLOWED", "不支持日期格式，请改为文本")
        )
        return ""
    if isinstance(value, float):
        if not math.isfinite(value):
            issues.append(
                _issue(row_number, column, "INVALID_NUMBER", "数字不是有效的有限值")
            )
            return ""
        if column in IDENTIFIER_COLUMNS and not value.is_integer():
            issues.append(
                _issue(row_number, column, "DECIMAL_IDENTIFIER", "标识字段不支持小数")
            )
            return ""
        value = int(value) if value.is_integer() else value

    text = str(value).strip()
    if len(text) > MAX_CELL_CHARACTERS:
        issues.append(
            _issue(
                row_number,
                column,
                "VALUE_TOO_LONG",
                f"内容超过 {MAX_CELL_CHARACTERS} 个字符",
            )
        )
    return text


def parse_dimension_workbook(
    content: bytes,
    original_name: str,
    dimension_file: str,
) -> dict:
    expected_columns = REQUIRED_DIMENSION_COLUMNS.get(dimension_file)
    if expected_columns is None:
        raise DimensionImportFileError("不支持的维表类型")
    if not original_name or Path(original_name).suffix.lower() != ".xlsx":
        raise DimensionImportFileError("仅支持 .xlsx 格式的 Excel 文件")
    if not content:
        raise DimensionImportFileError("Excel 文件不能为空")
    if len(content) > MAX_IMPORT_BYTES:
        raise DimensionImportFileError("Excel 文件不能超过 10 MB")

    _validate_archive(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise DimensionImportFileError("Excel 文件无法读取，请检查文件是否损坏") from exc

    issues: list[dict] = []
    rows: list[dict] = []
    try:
        if not workbook.worksheets:
            raise DimensionImportFileError("Excel 中没有可读取的工作表")
        worksheet = workbook.worksheets[0]
        # Some valid XLSX writers omit the optional worksheet dimension metadata.
        # In read-only mode openpyxl then exposes max_column as None even though
        # iter_rows can still read the cells normally.
        header_width = min(max(worksheet.max_column or 0, len(expected_columns)), 100)
        header_cells = next(
            worksheet.iter_rows(min_row=1, max_row=1, max_col=header_width),
            (),
        )
        actual_columns = [
            "" if cell.value is None else str(cell.value)
            for cell in header_cells
        ]
        while actual_columns and actual_columns[-1] == "":
            actual_columns.pop()

        headers_valid = actual_columns == expected_columns
        if not headers_valid:
            issues.append(
                _issue(
                    1,
                    "表头",
                    "HEADER_MISMATCH",
                    "表头名称和列顺序必须与当前维表完全一致",
                )
            )

        scan_width = max(len(actual_columns), len(expected_columns), 1)
        non_empty_rows = 0
        for row_number, cells in enumerate(
            worksheet.iter_rows(min_row=2, max_col=scan_width),
            start=2,
        ):
            raw_values = [cell.value for cell in cells]
            if all(value is None or str(value).strip() == "" for value in raw_values):
                continue
            non_empty_rows += 1
            if non_empty_rows > MAX_IMPORT_ROWS:
                issues.append(
                    _issue(
                        row_number,
                        "整行",
                        "TOO_MANY_ROWS",
                        f"单次最多导入 {MAX_IMPORT_ROWS:,} 行，请拆分文件",
                    )
                )
                break
            if not headers_valid:
                continue

            row_issues_before = len(issues)
            data = {
                column: _cell_text(cells[index], row_number, column, issues)
                for index, column in enumerate(expected_columns)
            }
            name_column = expected_columns[1]
            if not data.get("适用的包"):
                issues.append(
                    _issue(row_number, "适用的包", "REQUIRED", "适用的包不能为空")
                )
            if not data.get(name_column):
                issues.append(
                    _issue(row_number, name_column, "REQUIRED", f"{name_column}不能为空")
                )
            if len(issues) == row_issues_before:
                rows.append({"rowNumber": row_number, "data": data})

        if non_empty_rows == 0:
            issues.append(_issue(2, "整行", "NO_DATA", "Excel 中没有可导入的数据行"))

        return {
            "sourceName": original_name,
            "sheetName": worksheet.title,
            "rowCount": non_empty_rows,
            "columnCount": len(actual_columns),
            "expectedColumns": list(expected_columns),
            "actualColumns": actual_columns,
            "headersValid": headers_valid,
            "rows": rows,
            "issues": issues[:100],
            "errorCount": len(issues),
            "valid": not issues,
        }
    finally:
        workbook.close()
