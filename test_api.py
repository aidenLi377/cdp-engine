r"""
本地自测脚本。

运行方式:
    .venv\Scripts\python test_api.py
"""

from __future__ import annotations

import io
import json
import os
import unittest
from datetime import datetime

os.environ["FLASK_ENV"] = "development"

from cdp_backend.app_factory import is_production  # noqa: E402
from cdp_backend.validator import validate_project_config  # noqa: E402
from test_support import create_authenticated_test_app  # noqa: E402


class CdpApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_app = create_authenticated_test_app("api-test-user")
        cls.client = cls.test_app.client
        cls.engine = cls.test_app.engine

    @classmethod
    def tearDownClass(cls):
        cls.test_app.close()

    def test_environment(self):
        self.assertFalse(is_production())

    def test_config_validation(self):
        issues = validate_project_config()
        self.assertIsInstance(issues, list)

    def test_health(self):
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data["status"], "ok")

    def test_packages(self):
        response = self.client.get("/api/packages")
        packages = response.get_json()
        self.assertIsInstance(packages, list)
        self.assertIn("类目公域行为", packages)
        self.assertIn("商品行为", packages)
        self.assertIn("品牌专区", packages)
        self.assertIn("全媒体智投", packages)
        self.assertIn("单媒体智投", packages)
        self.assertIn("自定义人群", packages)

    def test_config_version_is_never_browser_cached(self):
        response = self.client.get("/api/config/version")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json(), {"version": 0, "publishedAt": None})
        self.assertIn("no-store", response.headers["Cache-Control"])

    def test_generate_custom_crowd_keeps_name_for_runtime_resolution(self):
        response = self.client.post(
            "/api/generate",
            json={"_package": "自定义人群", "crowdIds": "HN919I人群BH"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json(),
            {
                "crowdName": "未命名",
                "list": [{
                    "selectionLv1": ["CROWD", "CUSTOM"],
                    "selectionLv3": {"crowdIds": ["HN919I人群BH"]},
                    "fromPoolId": 0,
                }],
                "compute": "(0)",
            },
        )

    def test_meta_and_cache(self):
        response = self.client.get("/api/meta/类目公域行为?v=test-release.0")
        meta = response.get_json()
        self.assertIn("schema", meta)
        self.assertIn("matrix", meta)
        self.assertTrue(meta["schema"])
        self.assertIn("类目公域行为", self.engine._meta_cache)
        self.assertIn("max-age=31536000", response.headers["Cache-Control"])
        self.assertIn("immutable", response.headers["Cache-Control"])
        self.assertEqual(response.headers["Vary"], "Cookie")
        self.assertTrue(response.headers.get("ETag"))

        conditional = self.client.get(
            "/api/meta/类目公域行为?v=test-release.0",
            headers={"If-None-Match": response.headers["ETag"]},
        )
        self.assertEqual(conditional.status_code, 304)

        bundle_response = self.client.get("/api/meta?v=test-release.0")
        bundle = bundle_response.get_json()
        self.assertEqual(set(bundle), set(self.engine.packages))
        self.assertIn("schema", bundle["类目公域行为"])

        response_alias = self.client.get("/api/package_meta?name=类目公域行为")
        self.assertEqual(response_alias.status_code, 200)
        self.assertIn("no-cache", response_alias.headers["Cache-Control"])
        self.assertIn("max-age=0", response_alias.headers["Cache-Control"])
        alias_meta = response_alias.get_json()
        self.assertEqual(len(meta["schema"]), len(alias_meta["schema"]))

    def test_generate_category_json(self):
        payload = {
            "_package": "类目公域行为",
            "bhv": ["购买"],
            "leafCates": ["美容护肤/美体/精油>乳液/面霜"],
            "stdBrand": ["CPB/肌肤之钥"],
            "channel": ["天猫"],
            "frequency": {"min": "", "max": ""},
            "price": {"min": "", "max": ""},
            "itemprice": {"min": "", "max": ""},
            "time": {
                "val": {"start": "20260501", "end": "20260621"},
                "min": "range",
            },
        }
        response = self.client.post("/api/generate", json=payload)
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        expected = {
            "crowdName": "未命名",
            "list": [
                {
                    "selectionLv1": ["COMMON_TOUCH", "PUBLIC_CATE_BHV"],
                    "selectionLv3": {
                        "extraFilters": {
                            "channel": ["16772#|#4"],
                            "stdBrand": ["18641319"],
                            "frequency": {"op": "OPEN_OPEN"},
                            "price": {"op": "OPEN_OPEN"},
                            "itemprice": {"op": "OPEN_OPEN"},
                        },
                        "leafCates": ["50011980#|#50011980"],
                        "bhv": ["18919#|#CATE_PUBLIC_PAY"],
                        "dateType": "ABSOLUTE_DATE_RANGE",
                        "dateValue": {"from": "20260501", "to": "20260621"},
                    },
                    "fromPoolId": 0,
                }
            ],
            "compute": "(0)",
        }
        self.assertEqual(data, expected)
        self.assertEqual(list(data["list"][0]), ["selectionLv1", "selectionLv3", "fromPoolId"])
        self.assertEqual(
            list(data["list"][0]["selectionLv3"]),
            ["extraFilters", "leafCates", "bhv", "dateType", "dateValue"],
        )
        self.assertEqual(
            list(data["list"][0]["selectionLv3"]["extraFilters"]),
            ["channel", "stdBrand", "frequency", "price", "itemprice"],
        )
        response_text = response.get_data(as_text=True)
        self.assertLess(response_text.index('"selectionLv1"'), response_text.index('"selectionLv3"'))
        self.assertLess(response_text.index('"selectionLv3"'), response_text.index('"fromPoolId"'))

    def test_generate_category_json_without_optional_brand(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "类目公域行为",
                "bhv": ["购买"],
                "leafCates": ["美容护肤/美体/精油>乳液/面霜"],
                "channel": ["天猫"],
                "frequency": {"min": "", "max": ""},
                "price": {"min": "", "max": ""},
                "itemprice": {"min": "", "max": ""},
                "time": {
                    "val": {"start": "20260501", "end": "20260621"},
                    "min": "range",
                },
            },
        )

        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(len(data["list"]), 1)
        self.assertNotIn("stdBrand", data["list"][0]["selectionLv3"]["extraFilters"])

    def test_generate_category_json_normalizes_iso_date_range_from_ai_workbench(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "类目公域行为",
                "bhv": ["购买"],
                "leafCates": ["美容护肤/美体/精油>乳液/面霜"],
                "channel": ["天猫"],
                "frequency": {"min": "", "max": ""},
                "price": {"min": "", "max": ""},
                "itemprice": {"min": "", "max": ""},
                "time": {
                    "val": {"start": "2025-09-16", "end": "2026-03-14"},
                    "min": "range",
                },
            },
        )

        self.assertEqual(response.status_code, 200)
        date_value = response.get_json()["list"][0]["selectionLv3"]["dateValue"]
        self.assertEqual(date_value, {"from": "20250916", "to": "20260314"})

    def test_category_item_meta_supports_batch_ids(self):
        response = self.client.get("/api/meta/类目商品行为")
        self.assertEqual(response.status_code, 200)
        item_field = next(field for field in response.get_json()["schema"] if field["key"] == "item")
        self.assertEqual(item_field["Widget_Type"], "列表输入")
        self.assertIn("批量粘贴多个商品ID", item_field["Description"])

    def test_generate_category_item_json_keeps_single_item_id_scalar(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "类目商品行为",
                "bhv": ["购买"],
                "item": ["123456789"],
                "time": {"val": {"days": 30}, "min": "recent"},
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(
            data["list"][0]["selectionLv3"]["extraFilters"]["item"],
            "123456789",
        )

    def test_brand_zone_meta_uses_data_driven_single_choice_fields(self):
        response = self.client.get("/api/meta/品牌专区")
        self.assertEqual(response.status_code, 200)
        meta = response.get_json()
        fields = {field["key"]: field for field in meta["schema"]}

        self.assertEqual(list(fields), ["account", "bhv", "dayFrequency", "time"])
        self.assertEqual(fields["account"]["Widget_Type"], "单选组")
        self.assertEqual(
            fields["account"]["options"],
            ["以下全部投放账号", "dior迪奥官方旗舰店"],
        )
        self.assertEqual(fields["account"]["uiConfig"]["display"], "radio")
        self.assertEqual(fields["account"]["uiConfig"]["defaultValue"], "以下全部投放账号")
        self.assertEqual(fields["bhv"]["Widget_Type"], "单选组")
        self.assertEqual(fields["bhv"]["options"], ["被广告曝光过", "点击过广告"])
        self.assertEqual(fields["dayFrequency"]["uiConfig"]["minModeLabel"], "大于")
        self.assertEqual(fields["time"]["uiConfig"]["relativeModeLabel"], "相对日期")

    def test_generate_brand_zone_matches_official_examples(self):
        scenarios = [
            (
                "全部账号_曝光_天数区间_固定日期",
                {
                    "_package": "品牌专区",
                    "account": "以下全部投放账号",
                    "bhv": "被广告曝光过",
                    "dayFrequency": {"min": 1, "max": 10},
                    "time": {
                        "val": {"start": "20260901", "end": "20260910"},
                        "min": "range",
                    },
                },
                {
                    "contType": "bhv",
                    "account": "ALL",
                    "dayFrequency": {"op": "CLOSE_CLOSE", "min": 1, "max": 10},
                    "bhv": "15274#|#EXPOSE_AD",
                    "dateType": "ABSOLUTE_DATE_RANGE",
                    "dateValue": {"from": "20260901", "to": "20260910"},
                },
            ),
            (
                "全部账号_点击_天数不限_最近180天",
                {
                    "_package": "品牌专区",
                    "account": "以下全部投放账号",
                    "bhv": "点击过广告",
                    "dayFrequency": {"min": "", "max": ""},
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "contType": "bhv",
                    "account": "ALL",
                    "dayFrequency": {"op": "OPEN_OPEN"},
                    "bhv": "15300#|#CLICK_AD",
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                },
            ),
            (
                "指定账号_点击_大于180天_最近180天",
                {
                    "_package": "品牌专区",
                    "account": "dior迪奥官方旗舰店",
                    "bhv": "点击过广告",
                    "dayFrequency": {"min": 180, "max": ""},
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "contType": "bhv",
                    "account": "2207959261164#|#2207959261164",
                    "dayFrequency": {"op": "OPEN_CLOSE", "min": 180},
                    "bhv": "15300#|#CLICK_AD",
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                },
            ),
        ]

        for name, payload, expected_lv3 in scenarios:
            with self.subTest(name=name):
                response = self.client.post("/api/generate", json=payload)
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
                data = response.get_json()
                expected = {
                    "crowdName": "未命名",
                    "list": [
                        {
                            "selectionLv1": ["FIELD", "AD"],
                            "selectionLv3": expected_lv3,
                            "fromPoolId": 0,
                            "selectionLv2Name": "品牌专区",
                            "selectionLv2": ["15250#|#EB"],
                        }
                    ],
                    "compute": "(0)",
                }
                self.assertEqual(data, expected)
                self.assertEqual(
                    list(data["list"][0]),
                    ["selectionLv1", "selectionLv3", "fromPoolId", "selectionLv2Name", "selectionLv2"],
                )
                self.assertEqual(
                    list(data["list"][0]["selectionLv3"]),
                    ["contType", "account", "dayFrequency", "bhv", "dateType", "dateValue"],
                )

    def test_effect_promotion_meta_uses_behavior_specific_scenes(self):
        response = self.client.get("/api/meta/效果推广")
        self.assertEqual(response.status_code, 200)
        fields = {field["key"]: field for field in response.get_json()["schema"]}

        self.assertEqual(
            list(fields),
            ["account", "bhv", "onebp_scene", "dayFrequency", "time"],
        )
        self.assertEqual(fields["account"]["options"], ["全部", "dior迪奥官方旗舰店"])
        self.assertEqual(fields["bhv"]["options"], ["曝光", "点击", "观看"])
        self.assertIsNone(fields["account"]["uiConfig"]["defaultValue"])
        self.assertIsNone(fields["bhv"]["uiConfig"]["defaultValue"])
        self.assertEqual(fields["onebp_scene"]["Widget_Type"], "动态多选")
        self.assertEqual(
            {key: len(value) for key, value in fields["onebp_scene"]["optionsByValue"].items()},
            {"曝光": 17, "点击": 17, "观看": 3},
        )
        self.assertEqual(
            fields["onebp_scene"]["optionsByValue"]["观看"],
            ["超级直播", "超级短视频", "短直联动"],
        )
        self.assertEqual(
            fields["onebp_scene"]["uiConfig"]["displayByValue"]["观看"],
            "checkbox",
        )

    def test_generate_effect_promotion_matches_official_examples(self):
        exposure_scenes = [
            "货品运营",
            "关键词推广(原淘内广告/直通车)",
            "精准人群推广(整合原消费者运营)",
            "获客易",
            "其他场景推广-全店智投(原全店推)",
            "其他场景推广-活动加速",
            "其他场景推广-多目标直投",
            "其他场景推广-其他(对应原服务商场景)",
            "货品运营-测款快",
            "消费者运营-拉新快",
            "货品运营-货品加速",
            "货品运营-上新快",
            "消费者运营-会员快",
            "消费者运营-粉丝快",
            "消费者运营-追投快",
            "消费者运营-人群击穿",
            "原万相台-电商场景",
        ]
        exposure_values = [
            "19111#|#376", "15403#|#371", "15405#|#372", "15386#|#144",
            "15394#|#361", "15392#|#154", "19110#|#427", "15396#|#-100",
            "15348#|#158", "15374#|#78", "15369#|#114", "15363#|#105",
            "15378#|#133", "15390#|#407", "15388#|#189", "15382#|#370",
            "15401#|#-999",
        ]
        scenarios = [
            (
                "曝光_全选",
                "曝光",
                exposure_scenes,
                {
                    "account": "2207959261164#|#2207959261164",
                    "bhv": "15296#|#onebp_expose",
                    "onebp_scene": exposure_values,
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                    "dayFrequency": {"op": "OPEN_OPEN"},
                },
            ),
            (
                "点击_货品运营",
                "点击",
                ["货品运营"],
                {
                    "account": "2207959261164#|#2207959261164",
                    "bhv": "15316#|#onebp_click",
                    "onebp_scene": ["19113#|#376"],
                    "dateType": "RELATIVE_RANGE",
                    "dayFrequency": {"op": "OPEN_OPEN"},
                },
            ),
            (
                "观看_三个场景",
                "观看",
                ["超级直播", "超级短视频", "短直联动"],
                {
                    "account": "2207959261164#|#2207959261164",
                    "bhv": "15323#|#onebp_view",
                    "onebp_scene": ["15397#|#108", "15398#|#183", "15399#|#341"],
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                    "dayFrequency": {"op": "OPEN_OPEN"},
                },
            ),
        ]

        for name, behavior, scenes, expected_lv3 in scenarios:
            with self.subTest(name=name):
                response = self.client.post(
                    "/api/generate",
                    json={
                        "_package": "效果推广",
                        "account": "dior迪奥官方旗舰店",
                        "bhv": behavior,
                        "onebp_scene": scenes,
                        "dayFrequency": {"min": "", "max": ""},
                        "time": {"val": {"days": 180}, "min": "recent"},
                    },
                )
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
                self.assertEqual(
                    response.get_json(),
                    {
                        "crowdName": "未命名",
                        "list": [{
                            "selectionLv1": ["FIELD", "AD"],
                            "selectionLv3": expected_lv3,
                            "fromPoolId": 0,
                            "selectionLv2Name": "效果推广",
                            "selectionLv2": ["15270#|#cate_8954"],
                        }],
                        "compute": "(0)",
                    },
                )

    def test_effect_promotion_rejects_scene_from_another_behavior(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "效果推广",
                "account": "全部",
                "bhv": "观看",
                "onebp_scene": ["货品运营"],
                "dayFrequency": {"min": "", "max": ""},
                "time": {"val": {"days": 180}, "min": "recent"},
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["code"], "INVALID_PARAMETERS")
        self.assertIn("场景与观看行为不匹配", response.get_json()["message"])

    def test_brand_promotion_meta_uses_behavior_specific_multi_select_scenes(self):
        response = self.client.get("/api/meta/品牌推广")
        self.assertEqual(response.status_code, 200)
        fields = {field["key"]: field for field in response.get_json()["schema"]}

        self.assertEqual(
            list(fields),
            ["bhv", "ppob_scene", "cate", "dayFrequency", "time"],
        )
        self.assertEqual(fields["bhv"]["Widget_Type"], "单选组")
        self.assertEqual(fields["bhv"]["options"], ["点击", "曝光"])
        self.assertIsNone(fields["bhv"]["uiConfig"]["defaultValue"])
        self.assertEqual(fields["ppob_scene"]["Widget_Type"], "动态多选")
        self.assertEqual(
            {key: len(value) for key, value in fields["ppob_scene"]["optionsByValue"].items()},
            {"点击": 15, "曝光": 15},
        )
        self.assertEqual(fields["ppob_scene"]["optionsByValue"]["点击"][4], "一搜即现")
        self.assertEqual(fields["cate"]["Widget_Type"], "搜索单选")
        self.assertEqual(fields["cate"]["options"][0], "全部")
        self.assertIn("彩妆/香水/美妆工具>唇部彩妆", fields["cate"]["options"])
        self.assertEqual(fields["cate"]["uiConfig"]["defaultValue"], "全部")
        self.assertEqual(fields["dayFrequency"]["uiConfig"]["minModeLabel"], "大于")
        self.assertEqual(fields["time"]["uiConfig"]["defaultDays"], 180)

    def test_generate_brand_promotion_matches_supplied_examples(self):
        scenarios = [
            (
                "曝光_双场景_全部类目_不限_最近180天",
                {
                    "bhv": "曝光",
                    "ppob_scene": [
                        "UD全域品牌营销-超级亮相",
                        "淘内展示营销 - 人群击穿",
                    ],
                    "cate": "全部",
                    "dayFrequency": {"min": "", "max": ""},
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "cate": "ALL",
                    "bhv": "15318#|#exp_pptg",
                    "ppob_scene": ["20216#|#395", "20215#|#389"],
                    "dateType": "RELATIVE_RANGE",
                    "dayFrequency": {"op": "OPEN_OPEN"},
                },
            ),
            (
                "点击_一搜即现_全部类目_10至15天_最近180天",
                {
                    "bhv": "点击",
                    "ppob_scene": ["一搜即现"],
                    "cate": "全部",
                    "dayFrequency": {"min": 10, "max": 15},
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "cate": "ALL",
                    "bhv": "15298#|#click_pptg",
                    "ppob_scene": ["19487#|#353"],
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                    "dayFrequency": {"op": "CLOSE_CLOSE", "min": 10, "max": 15},
                },
            ),
            (
                "点击_一搜即现_唇部彩妆_大于180天_最近180天",
                {
                    "bhv": "点击",
                    "ppob_scene": ["一搜即现"],
                    "cate": "彩妆/香水/美妆工具>唇部彩妆",
                    "dayFrequency": {"min": 180, "max": ""},
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "cate": "201166702#|#201166702",
                    "bhv": "15298#|#click_pptg",
                    "ppob_scene": ["19487#|#353"],
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                    "dayFrequency": {"op": "OPEN_CLOSE", "min": 180},
                },
            ),
        ]

        for name, form_data, expected_lv3 in scenarios:
            with self.subTest(name=name):
                response = self.client.post(
                    "/api/generate",
                    json={"_package": "品牌推广", **form_data},
                )
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
                self.assertEqual(
                    response.get_json(),
                    {
                        "crowdName": "未命名",
                        "list": [{
                            "selectionLv1": ["FIELD", "AD"],
                            "selectionLv3": expected_lv3,
                            "fromPoolId": 1,
                            "selectionLv2Name": "品牌推广",
                            "selectionLv2": ["15272#|#cate_8969"],
                        }],
                        "compute": "(0)",
                    },
                )

    def test_brand_promotion_rejects_unknown_scene(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "品牌推广",
                "bhv": "点击",
                "ppob_scene": ["不存在的场景"],
                "cate": "全部",
                "dayFrequency": {"min": "", "max": ""},
                "time": {"val": {"days": 180}, "min": "recent"},
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["code"], "INVALID_PARAMETERS")
        self.assertIn("场景与点击行为不匹配", response.get_json()["message"])

    def test_omnimedia_meta_has_two_behaviors_and_expected_controls(self):
        response = self.client.get("/api/meta/全媒体智投")
        self.assertEqual(response.status_code, 200)
        fields = {field["key"]: field for field in response.get_json()["schema"]}

        self.assertEqual(list(fields), ["bhv", "dayFrequency", "time"])
        self.assertEqual(fields["bhv"]["Widget_Type"], "单选组")
        self.assertEqual(fields["bhv"]["options"], ["曝光", "点击"])
        self.assertEqual(fields["bhv"]["uiConfig"]["defaultValue"], "曝光")
        self.assertEqual(
            fields["bhv"]["uiConfig"]["packageNotice"],
            "原UD智汇投更名为全媒体智投",
        )
        self.assertEqual(fields["dayFrequency"]["uiConfig"]["minModeLabel"], "大于")
        self.assertEqual(fields["dayFrequency"]["uiConfig"]["rangeModeLabel"], "区间")
        self.assertEqual(fields["time"]["uiConfig"]["relativeModeLabel"], "相对日期")
        self.assertEqual(fields["time"]["uiConfig"]["absoluteModeLabel"], "固定日期")
        self.assertEqual(fields["time"]["uiConfig"]["defaultDays"], 180)

    def test_generate_omnimedia_matches_official_examples_and_modes(self):
        scenarios = [
            (
                "曝光_不限_最近180天",
                {
                    "bhv": "曝光",
                    "dayFrequency": {"min": "", "max": ""},
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "dayFrequency": {"op": "OPEN_OPEN"},
                    "bhv": "19873#|#EXP_UD_ZHT_EXP_BHV",
                    "bhv_type": "exp_udzht",
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                },
            ),
            (
                "点击_不限_最近180天",
                {
                    "bhv": "点击",
                    "dayFrequency": {"min": "", "max": ""},
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "dayFrequency": {"op": "OPEN_OPEN"},
                    "bhv": "19874#|#EXP_UD_ZHT_EXP_BHV",
                    "bhv_type": "click_udzht",
                    "dateType": "RELATIVE_RANGE",
                    "dateValue": "180",
                },
            ),
            (
                "曝光_大于2天_固定日期",
                {
                    "bhv": "曝光",
                    "dayFrequency": {"min": 2, "max": ""},
                    "time": {
                        "val": {"start": "2026-09-01", "end": "2026-09-10"},
                        "min": "range",
                    },
                },
                {
                    "dayFrequency": {"op": "OPEN_CLOSE", "min": 2},
                    "bhv": "19873#|#EXP_UD_ZHT_EXP_BHV",
                    "bhv_type": "exp_udzht",
                    "dateType": "ABSOLUTE_DATE_RANGE",
                    "dateValue": {"from": "20260901", "to": "20260910"},
                },
            ),
            (
                "点击_天数区间_固定日期",
                {
                    "bhv": "点击",
                    "dayFrequency": {"min": 2, "max": 5},
                    "time": {
                        "val": {"start": "20260901", "end": "20260910"},
                        "min": "range",
                    },
                },
                {
                    "dayFrequency": {"op": "CLOSE_CLOSE", "min": 2, "max": 5},
                    "bhv": "19874#|#EXP_UD_ZHT_EXP_BHV",
                    "bhv_type": "click_udzht",
                    "dateType": "ABSOLUTE_DATE_RANGE",
                    "dateValue": {"from": "20260901", "to": "20260910"},
                },
            ),
        ]

        for name, form_data, expected_lv3 in scenarios:
            with self.subTest(name=name):
                response = self.client.post(
                    "/api/generate",
                    json={"_package": "全媒体智投", **form_data},
                )
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
                data = response.get_json()
                self.assertEqual(
                    data,
                    {
                        "crowdName": "未命名",
                        "list": [{
                            "selectionLv1": ["FIELD", "AD"],
                            "selectionLv3": expected_lv3,
                            "tipProperty": {
                                "dateTo": "20240717",
                                "type": 3,
                                "dateFrom": "20240615",
                                "content": "原UD智汇投更名为全媒体智投",
                            },
                            "fromPoolId": 1,
                            "selectionLv2Name": "全媒体智投",
                            "selectionLv2": ["19872#|#EXP_UD_ZHT_EXP_BHV"],
                        }],
                        "compute": "(0)",
                    },
                )
                self.assertEqual(
                    list(data["list"][0]),
                    [
                        "selectionLv1",
                        "selectionLv3",
                        "tipProperty",
                        "fromPoolId",
                        "selectionLv2Name",
                        "selectionLv2",
                    ],
                )
                self.assertEqual(
                    list(data["list"][0]["selectionLv3"]),
                    ["dayFrequency", "bhv", "bhv_type", "dateType", "dateValue"],
                )

    def test_omnimedia_behavior_is_single_choice(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "全媒体智投",
                "bhv": ["曝光", "点击"],
                "dayFrequency": {"min": "", "max": ""},
                "time": {"val": {"days": 180}, "min": "recent"},
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["code"], "INVALID_PARAMETERS")
        self.assertIn("行为必须是曝光或点击", response.get_json()["message"])

    def test_single_media_meta_uses_only_behavior_and_time(self):
        response = self.client.get("/api/meta/单媒体智投")
        self.assertEqual(response.status_code, 200)
        fields = {field["key"]: field for field in response.get_json()["schema"]}

        self.assertEqual(list(fields), ["bhv", "time"])
        self.assertEqual(fields["bhv"]["Widget_Type"], "单选组")
        self.assertEqual(fields["bhv"]["options"], ["曝光", "点击"])
        self.assertEqual(fields["bhv"]["uiConfig"]["display"], "radio")
        self.assertEqual(fields["bhv"]["uiConfig"]["defaultValue"], "曝光")
        self.assertEqual(fields["time"]["uiConfig"]["relativeModeLabel"], "相对日期")
        self.assertEqual(fields["time"]["uiConfig"]["absoluteModeLabel"], "固定日期")
        self.assertEqual(fields["time"]["uiConfig"]["defaultDays"], 180)

    def test_generate_single_media_matches_supplied_examples(self):
        scenarios = [
            (
                "曝光_最近180天",
                {
                    "_package": "单媒体智投",
                    "bhv": "曝光",
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "bhv": "19937#|#is_pv_uddmt",
                    "dateType": "RELATIVE_RANGE",
                },
            ),
            (
                "点击_最近180天",
                {
                    "_package": "单媒体智投",
                    "bhv": "点击",
                    "time": {"val": {"days": 180}, "min": "recent"},
                },
                {
                    "bhv": "19938#|#is_click_uddmt",
                    "dateType": "RELATIVE_RANGE",
                },
            ),
            (
                "点击_固定日期",
                {
                    "_package": "单媒体智投",
                    "bhv": "点击",
                    "time": {
                        "val": {"start": "20260901", "end": "20260902"},
                        "min": "range",
                    },
                },
                {
                    "bhv": "19938#|#is_click_uddmt",
                    "dateType": "ABSOLUTE_DATE_RANGE",
                    "dateValue": {"from": "20260901", "to": "20260902"},
                },
            ),
        ]

        for name, payload, expected_lv3 in scenarios:
            with self.subTest(name=name):
                response = self.client.post("/api/generate", json=payload)
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
                self.assertEqual(
                    response.get_json(),
                    {
                        "crowdName": "未命名",
                        "list": [{
                            "selectionLv1": ["FIELD", "AD"],
                            "selectionLv3": expected_lv3,
                            "tipProperty": {
                                "dateTo": "20990129",
                                "type": 3,
                                "dateFrom": "20250814",
                            },
                            "fromPoolId": 0,
                            "selectionLv2Name": "单媒体智投",
                            "selectionLv2": ["19936#|#cate_9195"],
                        }],
                        "compute": "(0)",
                    },
                )

    def test_single_media_rejects_unknown_behavior(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "单媒体智投",
                "bhv": "观看",
                "time": {"val": {"days": 180}, "min": "recent"},
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["code"], "INVALID_PARAMETERS")
        self.assertIn("行为必须是曝光或点击", response.get_json()["message"])

    def test_generate_commodity_json(self):
        payload = {
            "_package": "商品行为",
            "channel": "天猫",
            "shop": "SHISEIDO资生堂官方旗舰店",
            "bhv": ["购买"],
            "frequency": {"min": "", "max": ""},
            "money": {"min": "", "max": ""},
            "cate": "全部",
            "time": {
                "val": {"start": "20260501", "end": "20260621"},
                "min": "range",
            },
            "selectedGoodsType": "任意品牌商品",
        }
        response = self.client.post("/api/generate", json=payload)
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        expected = {
            "crowdName": "未命名",
            "list": [
                {
                    "selectionLv1": ["COMMODITY", "ITEM"],
                    "selectionLv3": {
                        "shop": "113498758#|#113498758",
                        "keywords": None,
                        "cate": "ALL",
                        "bhv": ["16709#|#PAY"],
                        "frequency": {"op": "OPEN_OPEN"},
                        "money": {"op": "OPEN_OPEN"},
                        "dateType": "ABSOLUTE_DATE_RANGE",
                        "dateValue": {"from": "20260501", "to": "20260621"},
                        "selectedGoodsType": "1",
                    },
                    "tipProperty": None,
                    "fromPoolId": 0,
                    "selectionLv2": ["16612#|#4"],
                    "selectionLv2Name": "天猫",
                }
            ],
            "compute": "(0)",
        }
        self.assertEqual(data, expected)
        self.assertEqual(
            list(data["list"][0]),
            [
                "selectionLv1",
                "selectionLv3",
                "tipProperty",
                "fromPoolId",
                "selectionLv2",
                "selectionLv2Name",
            ],
        )
        self.assertEqual(
            list(data["list"][0]["selectionLv3"]),
            [
                "shop",
                "keywords",
                "cate",
                "bhv",
                "frequency",
                "money",
                "dateType",
                "dateValue",
                "selectedGoodsType",
            ],
        )
        response_text = response.get_data(as_text=True)
        self.assertLess(response_text.index('"selectionLv1"'), response_text.index('"selectionLv3"'))
        self.assertLess(response_text.index('"selectionLv3"'), response_text.index('"tipProperty"'))
        self.assertLess(response_text.index('"tipProperty"'), response_text.index('"fromPoolId"'))

    def test_generate_commodity_multiple_behaviors_use_channel_codes(self):
        for channel, expected in (
            ("天猫", ["16709#|#PAY", "16625#|#VIEW_ITEM"]),
            ("所有销售渠道", ["16712#|#PAY", "16628#|#VIEW_ITEM"]),
        ):
            with self.subTest(channel=channel):
                response = self.client.post(
                    "/api/generate",
                    json={
                        "_package": "商品行为",
                        "channel": channel,
                        "bhv": ["购买", "浏览"],
                    },
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    response.get_json()["list"][0]["selectionLv3"]["bhv"],
                    expected,
                )

    def test_generate_commodity_behavior_without_channel_fails(self):
        response = self.client.post(
            "/api/generate",
            json={"_package": "商品行为", "bhv": ["购买", "浏览"]},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["code"], "INVALID_PARAMETERS")
        self.assertIn("请选择一个渠道", response.get_json()["message"])

    def test_generate_commodity_behavior_without_mapping_fails(self):
        response = self.client.post(
            "/api/generate",
            json={
                "_package": "商品行为",
                "channel": "天猫",
                "bhv": ["未配置行为"],
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["code"], "INVALID_PARAMETERS")
        self.assertIn("未配置行为代码", response.get_json()["message"])

    def test_generate_alias_json(self):
        payload = {
            "pkgName": "预测购买力",
            "params": {"attributes": ["高购买力"]},
        }
        response = self.client.post("/api/generate_json", json=payload)
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(data["list"])

    def test_batch_generate(self):
        csv_content = "人群名称,属性值\n测试包1,高购买力\n测试包2,L2\n"
        data = {
            "file": (io.BytesIO(csv_content.encode("utf-8-sig")), "批量圈人_预测购买力_模版.csv")
        }
        response = self.client.post("/api/batch_generate", data=data, content_type="multipart/form-data")
        payload = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(payload["detected_pkg"], "预测购买力")
        self.assertEqual(len(payload["results"]), 2)
        self.assertEqual(payload["errors"], [])

    def test_list_templates(self):
        response = self.client.get("/api/list_templates")
        templates = response.get_json()
        self.assertIsInstance(templates, list)
        self.assertTrue(any(name.endswith(".csv") for name in templates))

    def test_export_audience_run_has_four_columns_and_preserves_special_counts(self):
        response = self.client.post(
            "/api/audience-runs/export",
            json={
                "rows": [
                    {
                        "crowdName": "新品兴趣人群0921",
                        "crowdCount": 0,
                        "countObtainedAt": "2026-09-22T07:08:09.000Z",
                        "parameters": '{"crowdName":"新品兴趣人群0921","list":[]}',
                    },
                    {
                        "crowdName": "高潜购买人群0921",
                        "crowdCount": 7257408,
                        "countObtainedAt": "2026-09-22T07:09:10+00:00",
                        "parameters": '{"compute":"(0)"}',
                    },
                    {
                        "crowdName": "低量级人群0921",
                        "crowdCount": "-",
                        "countObtainedAt": None,
                        "parameters": '{"compute":"(1)"}',
                    },
                    {
                        "crowdName": "隐私阈值人群0921",
                        "crowdCount": "< 2,000",
                        "parameters": '{"compute":"(2)"}',
                    },
                ],
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(response.data), data_only=True)
        sheet = workbook["人群包结果"]
        values = list(sheet.iter_rows(values_only=True))
        self.assertEqual(values[0], ("人群包名称", "对应人数", "获取人数时间", "人群包参数"))
        self.assertEqual(len(values[0]), 4)
        self.assertEqual(values[1][0], "新品兴趣人群0921")
        self.assertEqual(values[1][1], 0)
        self.assertEqual(values[2][1], 7257408)
        self.assertEqual(values[3][1], "-")
        self.assertEqual(values[4][1], "<2000")
        self.assertEqual(values[1][2], datetime(2026, 9, 22, 15, 8, 9))
        self.assertEqual(values[2][2], datetime(2026, 9, 22, 15, 9, 10))
        self.assertIsNone(values[3][2])
        self.assertNotIn("\n", values[1][3])
        self.assertEqual(json.loads(values[1][3])["crowdName"], "新品兴趣人群0921")
        self.assertEqual(sheet["C2"].number_format, "yyyy-mm-dd hh:mm:ss")
        self.assertFalse(sheet["D2"].alignment.wrap_text)
        self.assertEqual(sheet.row_dimensions[1].height, 20)
        self.assertEqual(sheet.row_dimensions[2].height, 18)
        workbook.close()

    def test_errors(self):
        response = self.client.get("/api/non-existent-endpoint")
        self.assertEqual(response.status_code, 404)
        response = self.client.get("/api/package_meta")
        self.assertEqual(response.status_code, 400)

    def test_route_interface_demo_page(self):
        response = self.client.get("/route-interface-demo")
        try:
            self.assertEqual(response.status_code, 200)
            self.assertIn("text/html", response.content_type)
            body = response.get_data(as_text=True)
            self.assertIn("Route vs API", body)
            self.assertIn("/api/solutions", body)
        finally:
            response.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
